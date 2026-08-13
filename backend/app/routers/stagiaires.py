from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Query
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from typing import Optional, List
from datetime import datetime, date
from io import BytesIO
import logging
import re
import unicodedata
import openpyxl
from pydantic import BaseModel
from app.database import get_db
from app.schemas.stagiaire import StagiaireCreate, StagiaireUpdate, StagiaireOut, StagiaireDetail
from app.services import stagiaire_service, encadrant_service
from app.auth.dependencies import get_current_user, require_role
from app.models.user import User, UserRole
from app.models.stagiaire import Stagiaire
from app.models.status import StatutStage, StatutDossier, get_allowed_stage_transitions, get_allowed_dossier_transitions, compute_statut_stage

logger = logging.getLogger("stagiaires.import")

router = APIRouter(prefix="/api/stagiaires", tags=["Stagiaires"])


class StatutChange(BaseModel):
    statut: str


class TransitionsOut(BaseModel):
    current_stage: str
    allowed_stage: list
    current_dossier: Optional[str] = None
    allowed_dossier: list = []


class BatchDelete(BaseModel):
    ids: List[int]


@router.get("/", response_model=List[StagiaireOut])
async def list_stagiaires(
    search: Optional[str] = None,
    statut: Optional[str] = None,
    statut_dossier: Optional[str] = None,
    statut_rapport: Optional[str] = None,
    has_attestation: Optional[bool] = None,
    periode: Optional[str] = None,
    encadrant_id: Optional[str] = None,
    date_min: Optional[date] = Query(None),
    date_max: Optional[date] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    statuts_stage = [s.strip() for s in statut.split(",") if s.strip()] if statut else None
    statuts_dossier = [s.strip() for s in statut_dossier.split(",") if s.strip()] if statut_dossier else None
    statuts_rapport = [s.strip() for s in statut_rapport.split(",") if s.strip()] if statut_rapport else None
    periodes = [p.strip() for p in periode.split(",") if p.strip()] if periode else None
    encadrant_ids = [int(i.strip()) for i in encadrant_id.split(",") if i.strip()] if encadrant_id else None

    if current_user.role == UserRole.ENCADRANT:
        encadrant_ids_filter = [current_user.id]
    else:
        encadrant_ids_filter = encadrant_ids

    return await stagiaire_service.get_stagiaires(
        db,
        search=search,
        statuts_stage=statuts_stage,
        statuts_dossier=statuts_dossier,
        statuts_rapport=statuts_rapport,
        encadrant_id=current_user.id if current_user.role == UserRole.ENCADRANT else None,
        has_attestation=has_attestation,
        periodes=periodes,
        encadrant_ids=encadrant_ids_filter if current_user.role != UserRole.ENCADRANT else None,
        date_min=date_min,
        date_max=date_max,
    )


@router.get("/export/excel")
async def export_stagiaires_excel(
    search: Optional[str] = None,
    statut: Optional[str] = None,
    statut_dossier: Optional[str] = None,
    statut_rapport: Optional[str] = None,
    has_attestation: Optional[bool] = None,
    periode: Optional[str] = None,
    encadrant_id: Optional[str] = None,
    date_min: Optional[date] = Query(None),
    date_max: Optional[date] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN)),
):
    statuts_stage = [s.strip() for s in statut.split(",") if s.strip()] if statut else None
    statuts_dossier = [s.strip() for s in statut_dossier.split(",") if s.strip()] if statut_dossier else None
    statuts_rapport = [s.strip() for s in statut_rapport.split(",") if s.strip()] if statut_rapport else None
    periodes = [p.strip() for p in periode.split(",") if p.strip()] if periode else None
    encadrant_ids = [int(i.strip()) for i in encadrant_id.split(",") if i.strip()] if encadrant_id else None

    if current_user.role == UserRole.ENCADRANT:
        encadrant_ids_filter = [current_user.id]
    else:
        encadrant_ids_filter = encadrant_ids

    stagiaires = await stagiaire_service.get_stagiaires(
        db,
        search=search,
        statuts_stage=statuts_stage,
        statuts_dossier=statuts_dossier,
        statuts_rapport=statuts_rapport,
        encadrant_id=current_user.id if current_user.role == UserRole.ENCADRANT else None,
        has_attestation=has_attestation,
        periodes=periodes,
        encadrant_ids=encadrant_ids_filter if current_user.role != UserRole.ENCADRANT else None,
        date_min=date_min,
        date_max=date_max,
    )

    import csv
    import io

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow([
        "Nom et prénom", "CIN", "Date de naissance", "Email", "Téléphone",
        "Date début", "Date fin", "Période", "Encadrant",
        "Lettre d'affectation",
        "Statut stage", "Statut dossier", "Prolongation",
    ])

    for s in stagiaires:
        writer.writerow([
            s.nom_complet or "",
            s.cin or "",
            s.date_naissance.strftime("%d/%m/%Y") if s.date_naissance else "",
            s.email or "",
            f"{s.country_code or ''} {s.telephone}" if s.telephone else "",
            s.date_debut_stage.strftime("%d/%m/%Y") if s.date_debut_stage else "",
            s.date_fin_stage.strftime("%d/%m/%Y") if s.date_fin_stage else "",
            s.periode_stage or "",
            s.encadrant_nom or "",
            "Oui" if s.lettre_affectation else "Non",
            s.statut_stage or "",
            s.statut_dossier or "",
            "Oui" if getattr(s, "is_extended", False) else "Non",
        ])

    csv_content = output.getvalue()

    import datetime
    filename = f"Stagiaires_{datetime.date.today().strftime('%Y-%m-%d')}.csv"
    return Response(
        content=csv_content.encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/{stagiaire_id}", response_model=StagiaireDetail)
async def get_stagiaire(stagiaire_id: int, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    result = await stagiaire_service.get_stagiaire_detail(db, stagiaire_id)
    if not result:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Stagiaire introuvable")
    return result


@router.post("/{stagiaire_id}/statut", response_model=StagiaireOut)
async def change_statut(stagiaire_id: int, data: StatutChange, db: AsyncSession = Depends(get_db), current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN))):
    try:
        StatutStage(data.statut)
        return await stagiaire_service.change_statut_stage(db, stagiaire_id, data.statut)
    except ValueError:
        try:
            StatutDossier(data.statut)
            return await stagiaire_service.change_statut_dossier(db, stagiaire_id, data.statut)
        except ValueError:
            raise HTTPException(status_code=400, detail=f"Statut invalide : {data.statut}")


@router.post("/{stagiaire_id}/statut-stage", response_model=StagiaireOut)
async def change_statut_stage(stagiaire_id: int, data: StatutChange, db: AsyncSession = Depends(get_db), current_user: User = Depends(require_role(UserRole.ENCADRANT, UserRole.RH, UserRole.ADMIN))):
    return await stagiaire_service.change_statut_stage(db, stagiaire_id, data.statut, current_user)


@router.post("/{stagiaire_id}/statut-dossier", response_model=StagiaireOut)
async def change_statut_dossier(stagiaire_id: int, data: StatutChange, db: AsyncSession = Depends(get_db), current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN))):
    return await stagiaire_service.change_statut_dossier(db, stagiaire_id, data.statut)


@router.get("/{stagiaire_id}/transitions", response_model=TransitionsOut)
async def get_transitions(stagiaire_id: int, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    result = await stagiaire_service.get_stagiaire_by_id(db, stagiaire_id)
    if not result:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Stagiaire introuvable")

    current_stage = result.statut_stage
    if isinstance(current_stage, str):
        try:
            current_stage = StatutStage(current_stage)
        except ValueError:
            current_stage = StatutStage.STAGE_NON_DEBUTE

    current_dossier = None
    if result.statut_dossier:
        try:
            current_dossier = StatutDossier(result.statut_dossier)
        except ValueError:
            pass

    allowed_stage = get_allowed_stage_transitions(current_stage)
    allowed_dossier = get_allowed_dossier_transitions(current_dossier)

    return TransitionsOut(
        current_stage=current_stage.value,
        allowed_stage=[s.value for s in allowed_stage],
        current_dossier=current_dossier.value if current_dossier else None,
        allowed_dossier=[s.value for s in allowed_dossier],
    )


@router.post("/", response_model=StagiaireOut)
async def create_stagiaire(data: StagiaireCreate, db: AsyncSession = Depends(get_db), current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN))):
    return await stagiaire_service.create_stagiaire(db, data)


@router.post("/batch-delete")
async def delete_stagiaires_batch(data: BatchDelete, db: AsyncSession = Depends(get_db), current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN))):
    if not data.ids:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Aucun stagiaire sélectionné")
    deleted = await stagiaire_service.delete_stagiaires(db, data.ids)
    return {"deleted": deleted}


@router.put("/{stagiaire_id}", response_model=StagiaireOut)
async def update_stagiaire(stagiaire_id: int, data: StagiaireUpdate, db: AsyncSession = Depends(get_db), current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN))):
    result = await stagiaire_service.update_stagiaire(db, stagiaire_id, data)
    if not result:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Stagiaire introuvable")
    return result


def _parse_excel_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        from datetime import timedelta
        return datetime(1899, 12, 30) + timedelta(days=value)
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y", "%m/%d/%Y", "%Y%m%d", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Format de date invalide : {value}")


COLUMN_PATTERNS = {
    "nom_complet": ["nom et prénom", "nom complet", "nom et prenom", "nom complet stagiaire", "full name", "nom complet du stagiaire"],
    "nom": ["nom", "name", "last name", "nom stagiaire", "family name", "surname"],
    "prenom": ["prénom", "prenom", "first name", "prénom stagiaire", "given name", "firstname"],
    "lettre_affectation": ["lettre d'affectation", "lettre affectation", "lettre d affectation", "affectation", "lettre", "assignment letter"],
    "date_naissance": ["date naissance", "date de naissance", "naissance", "birth date", "date of birth", "ddn", "birthday", "né", "nee", "né(e)"],
    "cin": ["cin", "c.i.n", "n° cin", "numéro cin", "id card", "identity", "carte identité", "carte d'identité", "identifiant", "piece identité"],
    "telephone": ["téléphone", "telephone", "tel", "phone", "mobile", "portable", "gsm", "tél", "numéro téléphone", "numero telephone"],
    "email": ["email", "e-mail", "mail", "courriel", "adresse email", "e mail"],
    "date_debut_stage": ["date début", "date debut", "date début stage", "date debut stage", "début stage", "debut stage", "start date", "date de début", "début", "debut"],
    "date_fin_stage": ["date fin", "date fin stage", "fin stage", "end date", "date de fin", "fin"],
    "civilite": ["civilité", "civilite", "mr", "mme", "mlle", "genre", "sexe", "salutation", "titre"],
    "ecole": ["école", "ecole", "school", "université", "universite", "university", "établissement", "etablissement", "institution", "formation"],
    "service": ["service", "department", "département", "departement", "direction"],
    "encadrant_nom": ["encadrant", "supervisor", "tuteur", "maître de stage", "maitre de stage", "responsable", "encadrant pédagogique", "encadrant professionnel"],
    "periode_stage": ["période", "periode", "session", "période stage", "periode stage"],
    "taille": ["taille", "size", "vêtement", "vetement", "tenue"],
    "pointure": ["pointure", "chaussure", "shoe size", "chaussant"],
}


def _normalize_header(s):
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.lower().strip()


def _detect_column(headers, field_name):
    patterns = [_normalize_header(p) for p in COLUMN_PATTERNS.get(field_name, [field_name])]
    normalized = [_normalize_header(h) for h in headers]
    for i, hl in enumerate(normalized):
        if not hl:
            continue
        for p in patterns:
            if hl == p or hl.startswith(p) or p.startswith(hl):
                return i
    for i, hl in enumerate(normalized):
        if not hl:
            continue
        for p in patterns:
            if p in hl or hl in p:
                return i
    return None


def _sheet_has_data(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                return True
    return False


def _row_has_data(row):
    return any(v is not None and str(v).strip() != "" for v in row)


def _find_header_row(ws):
    for r_idx in range(1, ws.max_row + 1):
        values = [str(c.value or "").strip() if c.value is not None else "" for c in ws[r_idx]]
        if not any(values):
            continue
        if (_detect_column(values, "nom_complet") is not None
                or _detect_column(values, "nom") is not None
                or _detect_column(values, "prenom") is not None):
            return r_idx
    return None


def _normalize_nom_complet(value):
    if not value:
        return ""
    return str(value).strip()


def _normalize_nom(value):
    if not value:
        return ""
    return str(value).strip().upper()


def _normalize_prenom(value):
    if not value:
        return ""
    s = str(value).strip().lower()
    return s[0].upper() + s[1:] if s else ""


def _parse_oui_non(value):
    if value is None:
        return False
    s = str(value).strip().lower()
    return s in ("oui", "o", "yes", "y", "true", "1", "vrai", "ok")


def _to_int(v):
    if v is None:
        return None
    try:
        return int(v)
    except (ValueError, TypeError):
        return None


def _extract_row_data(row, col_map):
    row_data = {}
    for field, idx in col_map.items():
        row_data[field] = row[idx] if idx < len(row) else None
    return row_data


async def _process_import_file(db, file: UploadFile):
    result = {
        "filename": file.filename,
        "success_count": 0,
        "errors": [],
        "status": "ok",
        "sheets": [],
    }
    if not file.filename.lower().endswith(".xlsx"):
        if file.filename.lower().endswith(".xls"):
            result["status"] = "ignored"
            result["errors"].append({"row": 0, "message": "Format .xls non supporté, utilisez .xlsx"})
            return result
        result["status"] = "ignored"
        result["errors"].append({"row": 0, "message": "Format non Excel (.xlsx)"})
        return result

    try:
        content = await file.read()
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
        if not any(_sheet_has_data(ws) for ws in wb.worksheets):
            wb = openpyxl.load_workbook(BytesIO(content), data_only=False)
    except Exception:
        result["status"] = "ignored"
        result["errors"].append({"row": 0, "message": "Fichier illisible ou corrompu"})
        return result

    logger.info("Import '%s' : %d feuille(s) détectée(s)", file.filename, len(wb.worksheets))

    for ws in wb.worksheets:
        sheet_name = ws.title

        if not _sheet_has_data(ws):
            logger.info("Feuille '%s' ignorée : complètement vide", sheet_name)
            continue

        header_row = _find_header_row(ws)
        if header_row is None:
            logger.info("Feuille '%s' ignorée : aucune ligne d'en-tête Nom/Prénom détectée", sheet_name)
            result["errors"].append({"row": 0, "message": f"Feuille '{sheet_name}' : colonnes Nom/Prénom introuvables"})
            continue

        headers = [str(c.value or "").strip() if c.value is not None else "" for c in ws[header_row]]

        col_map = {}
        for field in COLUMN_PATTERNS:
            idx = _detect_column(headers, field)
            if idx is not None:
                col_map[field] = idx

        if "nom_complet" not in col_map and "nom" not in col_map and "prenom" not in col_map:
            logger.info("Feuille '%s' ignorée : aucune colonne de nom détectée", sheet_name)
            result["errors"].append({"row": 0, "message": f"Feuille '{sheet_name}' : colonnes Nom/Prénom introuvables"})
            continue

        data_row_count = max(0, ws.max_row - header_row)
        logger.info(
            "Feuille '%s' : ligne d'en-tête ligne %d, %d ligne(s) de données, colonnes reconnues : %s",
            sheet_name, header_row, data_row_count, ", ".join(col_map) or "aucune",
        )

        has_date_naissance = "date_naissance" in col_map
        has_date_debut = "date_debut_stage" in col_map
        has_date_fin = "date_fin_stage" in col_map

        sheet_imported = 0
        sheet_skipped = 0

        for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not _row_has_data(row):
                continue
            try:
                row_data = _extract_row_data(row, col_map)

                nom_complet = _normalize_nom_complet(row_data.get("nom_complet", ""))
                if not nom_complet:
                    nom = _normalize_nom(row_data.get("nom", ""))
                    prenom = _normalize_prenom(row_data.get("prenom", ""))
                    nom_complet = " ".join(part for part in (nom, prenom) if part)

                row_label = nom_complet or f"Ligne {row_idx}"

                if not nom_complet:
                    sheet_skipped += 1
                    message = "Nom et prénom manquants (ligne ignorée)"
                    result["errors"].append({"row": row_idx, "message": message})
                    logger.info("Feuille '%s' ligne %d : %s", sheet_name, row_idx, message)
                    continue

                date_naissance = _parse_excel_date(row_data.get("date_naissance")) if has_date_naissance else None
                date_debut = _parse_excel_date(row_data.get("date_debut_stage")) if has_date_debut else None
                date_fin = _parse_excel_date(row_data.get("date_fin_stage")) if has_date_fin else None

                if date_fin and date_debut and date_fin <= date_debut:
                    logger.warning(
                        "Feuille '%s' ligne %d : dates incohérentes (fin <= début) mais stagiaire importé quand même",
                        sheet_name, row_idx,
                    )

                cin = str(row_data.get("cin", "") or "").strip()
                telephone = str(row_data.get("telephone", "") or "").strip()
                email = str(row_data.get("email", "") or "").strip()

                if cin:
                    existing = await db.execute(select(Stagiaire).where(Stagiaire.cin == cin))
                    if existing.scalar_one_or_none():
                        message = f"{row_label} : CIN {cin} déjà existant"
                        result["errors"].append({"row": row_idx, "message": message})
                        logger.info("Feuille '%s' ligne %d : ignorée (CIN %s déjà existant)", sheet_name, row_idx, cin)
                        continue

                if email:
                    existing_email = await db.execute(select(Stagiaire).where(Stagiaire.email == email))
                    if existing_email.scalar_one_or_none():
                        message = f"{row_label} : email {email} déjà existant"
                        result["errors"].append({"row": row_idx, "message": message})
                        logger.info("Feuille '%s' ligne %d : ignorée (email %s déjà existant)", sheet_name, row_idx, email)
                        continue

                if date_debut and date_fin and date_fin > date_debut:
                    statut = compute_statut_stage(date_debut, date_fin).value
                else:
                    statut = "Stage non débuté"

                encadrant_nom_value = str(row_data.get("encadrant_nom", "") or "").strip() or None
                encadrant_id_value = None
                if encadrant_nom_value:
                    try:
                        enc = await encadrant_service.get_or_create_encadrant(db, encadrant_nom_value)
                        if enc:
                            encadrant_id_value = enc.user_id
                    except Exception as e:
                        logger.warning(
                            "Feuille '%s' ligne %d : impossible d'associer l'encadrant '%s' (%s)",
                            sheet_name, row_idx, encadrant_nom_value, e,
                        )

                stagiaire = Stagiaire(
                    nom_complet=nom_complet or "Inconnu",
                    lettre_affectation=_parse_oui_non(row_data.get("lettre_affectation")),
                    civilite=str(row_data.get("civilite", "") or "").strip() or None,
                    service=str(row_data.get("service", "") or "").strip() or None,
                    date_naissance=date_naissance,
                    cin=cin or f"TEMP-{row_idx}-{abs(hash(file.filename)) % 10000}",
                    country_code="+216",
                    telephone=telephone or "00000000",
                    email=email or f"temp_{row_idx}_{abs(hash(file.filename)) % 10000}@placeholder.fr",
                    ecole=str(row_data.get("ecole", "") or "").strip() or None,
                    taille=_to_int(row_data.get("taille")),
                    pointure=_to_int(row_data.get("pointure")),
                    date_debut_stage=date_debut,
                    date_fin_stage=date_fin,
                    periode_stage=str(row_data.get("periode_stage", "") or "").strip() or None,
                    statut_stage=statut,
                    statut_dossier=None,
                    encadrant_nom=encadrant_nom_value,
                    encadrant_id=encadrant_id_value,
                )
                db.add(stagiaire)
                async with db.begin_nested():
                    await db.flush()
                result["success_count"] += 1
                sheet_imported += 1

            except Exception as e:
                message = str(e)
                if "row_label" in locals() and row_label:
                    message = f"{row_label} : {message}"
                result["errors"].append({"row": row_idx, "message": message})
                logger.warning("Feuille '%s' ligne %d : erreur (%s)", sheet_name, row_idx, e)

        logger.info(
            "Feuille '%s' : %d stagiaire(s) importé(s), %d ligne(s) ignorée(s), %d erreur(s)",
            sheet_name, sheet_imported, sheet_skipped, len([e for e in result["errors"] if e["row"] != 0]),
        )
        result["sheets"].append({"name": sheet_name, "imported": sheet_imported})

    if result["success_count"] == 0 and not any(e["row"] == 0 for e in result["errors"]):
        result["status"] = "ignored"
        result["errors"].append({"row": 0, "message": "Aucune ligne de stagiaire valide trouvée dans le fichier"})

    logger.info("Import '%s' terminé : %d stagiaire(s) importé(s)", file.filename, result["success_count"])
    return result


@router.post("/import")
async def import_stagiaires(
    files: List[UploadFile] = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN)),
):
    total_files = len(files)
    total_imported = 0
    total_ignored = 0
    all_file_results = []
    all_errors = []

    for file in files:
        file_result = await _process_import_file(db, file)
        all_file_results.append(file_result)
        if file_result["status"] == "ignored" or (file_result["success_count"] == 0 and file_result["errors"]):
            total_ignored += 1
        elif file_result["success_count"] > 0:
            total_imported += 1
        for err in file_result["errors"]:
            all_errors.append({
                "file": file_result["filename"],
                "row": err["row"],
                "message": err["message"],
            })

    if total_imported > 0:
        await db.commit()

    total_stagiaires = sum(f["success_count"] for f in all_file_results)
    logger.info(
        "Import terminé : %d stagiaire(s) importé(s) sur %d fichier(s)",
        total_stagiaires, len(files),
    )

    return {
        "files_total": total_files,
        "files_imported": total_imported,
        "files_ignored": total_ignored,
        "stagiaires_imported": total_stagiaires,
        "errors_total": len(all_errors),
        "errors": all_errors,
    }


@router.delete("/{stagiaire_id}")
async def delete_stagiaire(stagiaire_id: int, db: AsyncSession = Depends(get_db), current_user: User = Depends(require_role(UserRole.RH, UserRole.ADMIN))):
    deleted = await stagiaire_service.delete_stagiaire(db, stagiaire_id)
    if not deleted:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Stagiaire introuvable")
    return {"detail": "Stagiaire supprime avec succes"}
